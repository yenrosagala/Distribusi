"""Smoke test untuk builder deck PPTX dashboard (tanpa kaleido)."""
import io

import pandas as pd

from modules.dashboard_page import _build_pptx


class FakeFig:
    def to_image(self, format="png", width=None, height=None):
        # PNG 1x1 transparan.
        return bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
            "890000000a49444154789c6360000002000154a24f4f0000000049454e44ae426082"
        )


def _section(moda="Laut"):
    return {
        "moda": moda, "bln": "Agustus", "thn": 2026,
        "figs": [FakeFig(), FakeFig()],
        "tables": [
            pd.DataFrame({"PERKEMBANGAN X (%)": [1.5, -2.0, None]},
                         index=["Papua", "Papua Barat", "Papua Tengah"]),
            pd.DataFrame({"PERKEMBANGAN Y (%)": [0.0, 12.34, -5.5]},
                         index=["Papua", "Papua Barat", "Papua Tengah"]),
        ],
        "narration": ("Total penumpang naik 5,00 persen.", "Kumulatif Y-on-Y menurun."),
    }


def test_build_pptx_smoke():
    data = _build_pptx([_section("Laut"), _section("Udara")])
    assert data[:2] == b"PK"  # zip magic pptx

    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 4  # 2 moda x 2 slide


def test_build_pptx_contains_content():
    from pptx import Presentation

    prs = Presentation(io.BytesIO(_build_pptx([_section()])))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        for shape in slide.shapes:
            if shape.has_table:
                table = shape.table
                texts.append(str(table.cell(0, 0).text))
                for r in range(1, len(table.rows)):
                    for c in range(1, len(table.columns)):
                        texts.append(str(table.cell(r, c).text))
    joined = "\n".join(texts)
    assert "Transportasi Laut" in joined
    assert "Agustus 2026" in joined
    assert "PROVINSI" in joined
    assert "1,50" in joined  # format Indonesia
    assert "-2,00" in joined
    assert "naik 5,00 persen" in joined


def test_empty_section_none_guard():
    assert _build_pptx([])[:2] == b"PK"