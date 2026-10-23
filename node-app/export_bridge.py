"""Render PowerPoint / Word / InDesign exports for the Node dashboard.

The Node app owns the HTTP surface; this script owns the document formats. It
calls the Streamlit project's own builders instead of re-implementing OOXML or
IDML in JavaScript, so both apps produce identical files.

Usage:  python export_bridge.py <kind> <payload.json> <out-file>
kinds:  pptx | docx | idml-transport | idml-pariwisata
"""

import io
import json
import logging
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

logging.disable(logging.WARNING)


def _chart_png(chart, width=1100, height=400):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    labels = chart["labels"]
    datasets = chart["datasets"]
    x = list(range(len(labels)))
    fig, ax = plt.subplots(figsize=(width / 100, height / 100), dpi=100)
    bottom = [0.0] * len(labels)
    for ds in datasets:
        y = [0.0 if v is None else float(v) for v in ds["data"]]
        ax.bar(x, y, bottom=bottom, width=0.62, color=ds.get("color") or "#1F4E79", label=ds["label"])
        bottom = [b + v for b, v in zip(bottom, y)]
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=10, color="#333333")
    if chart.get("title"):
        ax.set_title(chart["title"], fontsize=13, color="#0F172A")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.08), ncol=max(1, len(datasets)), frameon=False)
    ax.yaxis.grid(True, color="#D9D9D9", linewidth=0.8)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.tight_layout(pad=0.6)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def build_pptx(payload):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def title_slide(slide, text):
        box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.7))
        frame = box.text_frame
        frame.text = text
        run = frame.paragraphs[0].runs[0]
        run.font.size = Pt(24)
        run.font.bold = True

    def add_table(slide, table, left, top, width, height):
        columns = table["columns"]
        rows = table["rows"]
        shape = slide.shapes.add_table(len(rows) + 1, len(columns), left, top, width, height)
        grid = shape.table
        for j, head in enumerate(columns):
            grid.cell(0, j).text = str(head)
        for i, row in enumerate(rows, start=1):
            for j, value in enumerate(row):
                grid.cell(i, j).text = "" if value is None else str(value)
        for r_i, row in enumerate(grid.rows):
            row.height = Inches(0.3)
            for cell in row.cells:
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
                        if r_i == 0:
                            run.font.bold = True

    def add_narration(slide, paras, left, top, width, height):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        first = True
        for text in paras:
            para = frame.paragraphs[0] if first else frame.add_paragraph()
            para.text = text or ""
            first = False
            for run in para.runs:
                run.font.size = Pt(11)

    for sec in payload["sections"]:
        head = f"Transportasi {sec['moda']} — {sec['bln']} {sec['thn']}"

        s1 = prs.slides.add_slide(prs.slide_layouts[6])
        title_slide(s1, f"{head}: Grafik")
        x = Inches(0.4)
        for chart in sec.get("charts", []):
            png = _chart_png(chart)
            s1.shapes.add_picture(io.BytesIO(png), x, Inches(1.1), width=Inches(6.15))
            x += Inches(6.15) + Inches(0.25)

        s2 = prs.slides.add_slide(prs.slide_layouts[6])
        title_slide(s2, f"{head}: Tabel & Narasi")
        tables = sec.get("tables", [])
        if len(tables) > 0:
            add_table(s2, tables[0], Inches(0.4), Inches(1.1), Inches(6.15), Inches(2.2))
        if len(tables) > 1:
            add_table(s2, tables[1], Inches(6.8), Inches(1.1), Inches(6.15), Inches(2.2))
        add_narration(s2, sec.get("narration", []), Inches(0.4), Inches(3.6), Inches(12.5), Inches(3.6))

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _report_items(prov, thn, bln):
    from modules.report_page import collect_report_tables, generate_narrative_fallback, normalize_narrative

    items = collect_report_tables(prov, thn, bln)
    for item in items:
        region_label = "Bandara" if item["moda"] == "Transportasi Udara" else "Pelabuhan/Kabupaten"
        p1, p2 = generate_narrative_fallback(
            report_flat=item["report_flat"], col_target=item["col_target"], moda=item["moda"],
            region_label=region_label, bln=item["bln"], thn=item["thn"],
            prev_bln=item["prev_bln"], prev_thn=item["prev_thn"],
            col_prev=item["col_prev"], col_curr=item["col_curr"],
            col_cum_prev=item["col_cum_prev"], col_cum_curr=item["col_cum_curr"], prov=item["prov"],
        )
        item["p1"] = normalize_narrative(p1, item["col_target"])
        item["p2"] = normalize_narrative(p2, item["col_target"])
    return items


def build_docx(payload):
    from modules.report_page import create_complete_master_word_report

    prov, thn, bln = payload["prov"], int(payload["thn"]), payload["bln"]
    items = _report_items(prov, thn, bln)
    if not items:
        raise SystemExit(f"Tidak ada data untuk {prov} {bln} {thn}.")
    stream = create_complete_master_word_report(prov, thn, bln, items)
    return stream.getvalue()


def _asset(payload, key):
    import base64

    encoded = payload.get(f"{key}_b64")
    if not encoded:
        return None, None
    return base64.b64decode(encoded), payload.get(f"{key}_name") or f"{key}.png"


def build_idml_transport(payload):
    from modules.indesign_export import DEFAULT_TEMPLATE, fill_brs_template, paket_zip
    from modules.report_page import _kab_lookup_brs, get_comparison_data

    prov, thn, bln = payload["prov"], int(payload["thn"]), payload["bln"]
    items = _report_items(prov, thn, bln)
    if not items:
        raise SystemExit(f"Tidak ada data untuk {prov} {bln} {thn}.")

    laut_ports = None
    frame, *_ = get_comparison_data(prov, thn, bln, "Transportasi Laut")
    if frame is not None and not frame.empty:
        import re

        laut_ports = [re.split(r"\s*/\s*", str(v))[0] for v in frame["nama_pelabuhan"].drop_duplicates()]

    result = fill_brs_template(
        DEFAULT_TEMPLATE.read_bytes(), prov, thn, bln, items, _kab_lookup_brs,
        nomor_brs=(payload.get("nomor") or "").strip() or None,
        tanggal_rilis=(payload.get("tanggal") or "").strip() or None,
        laut_ports=laut_ports,
        qr_bytes=_asset(payload, "qr")[0],
        qr_name=_asset(payload, "qr")[1],
        infografis_bytes=_asset(payload, "infografis")[0],
        infografis_name=_asset(payload, "infografis")[1],
        template_dir=DEFAULT_TEMPLATE.parent,
    )
    name = f"BRS_Transportasi_{prov.replace(' ', '_')}_{bln}_{thn}.idml"
    return paket_zip(result.idml_bytes, name, template_dir=DEFAULT_TEMPLATE.parent)


def build_idml_pariwisata(payload):
    from modules.database import get_engine
    from modules.indesign_export import paket_zip
    from modules.indesign_pariwisata import fill_brs_pariwisata_template, kumpulkan_kelas, resolve_template

    prov, thn, bln = payload["prov"], int(payload["thn"]), payload["bln"]
    template_path = resolve_template(level=prov)
    if not Path(template_path).exists():
        raise SystemExit(f"Template InDesign pariwisata tidak ditemukan: {template_path}")
    data = kumpulkan_kelas(get_engine(), prov, thn, bln)
    if not data:
        raise SystemExit(f"Tidak ada data pariwisata untuk {prov} {bln} {thn}.")
    result = fill_brs_pariwisata_template(
        None, prov, thn, bln, data, level=prov,
        nomor_brs=(payload.get("nomor") or "").strip() or None,
        tanggal_rilis=(payload.get("tanggal") or "").strip() or None,
        qr_bytes=_asset(payload, "qr")[0],
        qr_name=_asset(payload, "qr")[1],
        infografis_bytes=_asset(payload, "infografis")[0],
        infografis_name=_asset(payload, "infografis")[1],
    )
    name = f"BRS_Pariwisata_{prov.replace(' ', '_')}_{bln}_{thn}.idml"
    return paket_zip(result.idml_bytes, name, template_dir=Path(template_path).parent)


BUILDERS = {
    "pptx": build_pptx,
    "docx": build_docx,
    "idml-transport": build_idml_transport,
    "idml-pariwisata": build_idml_pariwisata,
}


def main(argv):
    if len(argv) != 4:
        raise SystemExit(__doc__)
    kind, payload_path, out_path = argv[1], argv[2], argv[3]
    builder = BUILDERS.get(kind)
    if not builder:
        raise SystemExit(f"Jenis ekspor tidak dikenal: {kind}")
    with open(payload_path, "r", encoding="utf-8") as handle:
        payload = json.load(handle)
    data = builder(payload)
    with open(out_path, "wb") as handle:
        handle.write(data)
    print(f"{kind} {len(data)} bytes", file=sys.stderr)


if __name__ == "__main__":
    main(sys.argv)
